#!/usr/bin/env python3
"""将线索 Excel 转为浏览器可直接加载的静态数据文件。"""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_HEADERS = [
    "originalId",
    "areaName",
    "areaCode",
    "warAreaName",
    "warAreaCode",
    "cityName",
    "cityCode",
    "provinceName",
    "provinceCode",
    "dealerCode",
    "dealerShortname",
    "consultantName",
    "realName",
    "mobile",
    "intentLevel",
    "intentSeriesName",
    "companyName",
    "clueStatus",
    "periodTd",
    "sourceTypeOneName",
    "sourceTypeTwoName",
    "sourceTypeThreeName",
    "sourceTypeFourName",
    "audioCount",
    "audioExistStatus",
    "qcClueStatus",
    "finalStatus",
    "finalInvalidReason",
]
EXPECTED_SOURCE_COUNT = 2_000
EXPECTED_RETAINED_COUNT = 1_840


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def integer(value: object) -> int:
    try:
        return max(0, int(value or 0))
    except (TypeError, ValueError):
        return 0


def should_keep(row: dict[str, object]) -> bool:
    return bool(
        text(row["realName"])
        and text(row["intentLevel"])
        and text(row["consultantName"]) != "-"
    )


def normalize(row: dict[str, object]) -> dict[str, object]:
    recording_count = integer(row["audioCount"])
    return {
        "id": text(row["originalId"]),
        "brand": text(row["companyName"]),
        "province": text(row["provinceName"]),
        "provinceCode": text(row["provinceCode"]),
        "city": text(row["cityName"]),
        "cityCode": text(row["cityCode"]),
        "dealerCode": text(row["dealerCode"]),
        "store": text(row["dealerShortname"]),
        "region": text(row["areaName"]),
        "regionCode": text(row["areaCode"]),
        "zone": text(row["warAreaName"]),
        "zoneCode": text(row["warAreaCode"]),
        "advisorName": text(row["consultantName"]),
        "customerName": text(row["realName"]),
        "customerPhone": text(row["mobile"]),
        "intentGrade": text(row["intentLevel"]),
        "carSeries": text(row["intentSeriesName"]),
        "leadStatus": text(row["clueStatus"]),
        "leadDate": text(row["periodTd"]),
        "leadSource": text(row["sourceTypeOneName"]),
        "secondSource": text(row["sourceTypeTwoName"]),
        "thirdSource": text(row["sourceTypeThreeName"]),
        "fourthSource": text(row["sourceTypeFourName"]),
        "recordingCount": recording_count,
        "validRecording": "是" if integer(row["audioExistStatus"]) else "否",
        "qcClueStatus": text(row["qcClueStatus"]),
        "finalStatus": text(row["finalStatus"]),
        "finalInvalidReason": text(row["finalInvalidReason"]),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="线索数据.xlsx 路径")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "leads" / "real-lead-data.js",
        help="生成的 JavaScript 数据文件",
    )
    args = parser.parse_args()

    workbook = load_workbook(args.source, read_only=True, data_only=True)
    if "Sheet1" not in workbook.sheetnames:
        raise ValueError("缺少工作表 Sheet1")
    worksheet = workbook["Sheet1"]
    rows = worksheet.iter_rows(values_only=True)
    headers = [text(value) for value in next(rows)]
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"Excel 字段不符合预期：{headers}")

    source_rows = [dict(zip(headers, values)) for values in rows]
    records = [normalize(row) for row in source_rows if should_keep(row)]
    ids = [record["id"] for record in records]
    if len(source_rows) != EXPECTED_SOURCE_COUNT:
        raise ValueError(
            f"源数据数量异常：期望 {EXPECTED_SOURCE_COUNT} 条，实际 {len(source_rows)} 条"
        )
    if len(records) != EXPECTED_RETAINED_COUNT:
        raise ValueError(
            f"清洗结果异常：期望 {EXPECTED_RETAINED_COUNT} 条，实际 {len(records)} 条"
        )
    invalid_ids = [value for value in ids if not value.isdigit() or len(value) != 19]
    if invalid_ids:
        raise ValueError(f"存在非 19 位数字线索 ID：{invalid_ids[:3]}")
    if len(ids) != len(set(ids)):
        raise ValueError("清洗后的线索 ID 存在重复")

    source_hash = hashlib.sha256(args.source.read_bytes()).hexdigest()
    payload = json.dumps(records, ensure_ascii=False, separators=(",", ":"))
    content = (
        "/* 由 scripts/generate-leads-real-data.py 生成，请勿手工修改。\n"
        f" * source: {args.source.name}; sha256: {source_hash}\n"
        f" * source rows: {len(source_rows)}; retained rows: {len(records)}\n"
        " */\n"
        f"window.__LEADS_REAL_DATA = Object.freeze({payload});\n"
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(content, encoding="utf-8")
    print(
        f"generated {args.output}: source={len(source_rows)}, "
        f"excluded={len(source_rows) - len(records)}, retained={len(records)}"
    )


if __name__ == "__main__":
    main()
